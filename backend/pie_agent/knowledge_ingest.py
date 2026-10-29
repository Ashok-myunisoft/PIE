"""v0.19 knowledge ingestion into the Qdrant KMS.

Two sources:
  * Usage catalogs (v0.19): module, capability and feature knowledge with verified
    and review claims, so the Marketing chat retrieves evidence-backed knowledge.
  * Tickets: support tickets from an export file (CSV / Excel / JSON) or from a
    read-only SQL SELECT. Personal data (reporter names, e-mails, phone numbers) is
    not stored.

All point IDs are deterministic, so re-running an import updates points in place.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import uuid
from datetime import datetime, timezone, date
from pathlib import Path
from typing import Any, Iterable

from .qdrant_kms import QdrantKMS

INGEST_VERSION = "0.19.0"
PRODUCT = "GoodBooks ERP"


def _uuid(kind: str, *parts: str) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"pie://{kind}/" + "/".join(parts)))


def with_vectors(points: list[dict[str, Any]], kms: QdrantKMS) -> list[dict[str, Any]]:
    """Points built with `text` get their vectors in batched embedding calls (one per ~256 texts)."""
    from .embeddings import prime
    prime(kms.embedder, [p["text"] for p in points if "text" in p])
    for p in points:
        if "text" in p:
            p["vector"] = kms.embedder.embed(p.pop("text"))
    return points


def _base_payload(**kw: Any) -> dict[str, Any]:
    payload = {
        "product": PRODUCT, "module": "", "functionality": "", "capability": "", "capability_basis": "",
        "capability_needs_review": False, "feature": "", "screen_id": "", "screen_name": "",
        "knowledge_type": "", "title": "", "content": "", "evidence_ids": [], "source_types": [],
        "confidence": 1.0, "review_required": False, "knowledge_review_required": False,
        "publication_status": "PUBLISHED", "generation_method": "deterministic", "claim_type": "observed",
        "source_hash": "", "knowledge_version": INGEST_VERSION,
        "generated_at": datetime.now(timezone.utc).isoformat(),
    }
    payload.update(kw)
    return payload


# ---------------------------------------------------------------- usage catalog

CONTEXT_KEYS = ("functionality", "capability", "capability_basis", "capability_needs_review")


def screen_contexts(module_dir: str | Path) -> dict[str, dict[str, Any]]:
    """Business context of every screen of a module, for knowledge built per screen (fields, sections):
    functionality from the screen registry (Phase 13) and capability from the Usage catalog (Phase 19).
    Screens missing from either source simply get fewer keys."""
    from .paths import from_root
    mdir = from_root(module_dir)
    ctx: dict[str, dict[str, Any]] = {}

    def read(rel: str) -> Any:
        f = mdir / rel
        return json.loads(f.read_text(encoding="utf-8")) if f.exists() else None
    for s in read("phase13-identity/screen_registry_v13.json") or []:
        if s.get("screenId") and s.get("functionality"):
            ctx.setdefault(s["screenId"], {})["functionality"] = s["functionality"]
    for f in (read("phase19-usage/usage_catalog_v19.json") or {}).get("features", []):
        cap = f.get("capability") or {}
        if not cap.get("name"):
            continue
        c = {"capability": cap["name"], "capability_basis": cap.get("basis") or "", "capability_needs_review": bool(cap.get("needsReview"))}
        for sid in [f["screenId"], *(f.get("aliasScreenIds") or [])]:
            ctx.setdefault(sid, {}).update(c)
    return ctx


def _context(ctx: dict[str, Any] | None) -> dict[str, Any]:
    return {k: ctx[k] for k in CONTEXT_KEYS if ctx and k in ctx}


def usage_catalog_points(catalog: dict[str, Any], kms: QdrantKMS, functionality: dict[str, str] | None = None) -> list[dict[str, Any]]:
    """Module overview + one point per capability + one per feature. `functionality`: screen id -> functionality."""
    module = catalog["module"]
    func = functionality or {}
    join = lambda sids: ", ".join(sorted({func[s] for s in sids if func.get(s)}))
    points: list[dict[str, Any]] = []
    stats = catalog["statistics"]

    caps_text = "; ".join(f"{c['name']} ({c['features']} features, {c['ready']} marketing-ready)" for c in catalog["capabilities"])
    content = (f"{catalog['moduleName']} module of {PRODUCT}: {stats['features']} features in {stats['capabilities']} capabilities, "
               f"{stats['ready']} marketing-ready, {stats['verifiedClaims']} verified highlights. Capabilities: {caps_text}.")
    points.append({"id": _uuid("usage", module, "module", INGEST_VERSION), "text": content,
                   "payload": _base_payload(knowledge_id=f"USG-{module}-MODULE", module=module, knowledge_type="UsageModule",
                                            functionality=join(f["screenId"] for f in catalog["features"]),
                                            title=f"{catalog['moduleName']} module overview", content=content,
                                            source_types=["Usage Catalog"], statistics={k: v for k, v in stats.items() if k != "sharedPlatformTables"},
                                            capabilities=[c["name"] for c in catalog["capabilities"]])})

    by_cap: dict[str, list[dict[str, Any]]] = {}
    for f in catalog["features"]:
        by_cap.setdefault(f["capability"]["name"], []).append(f)
    for cap in catalog["capabilities"]:
        feats = by_cap.get(cap["name"], [])
        names = ", ".join(f["name"] for f in feats)
        content = (f"Capability {cap['name']} in the {catalog['moduleName']} module: {cap['features']} features "
                   f"({cap['ready']} marketing-ready). Features: {names}.")
        points.append({"id": _uuid("usage", module, "capability", cap["name"], INGEST_VERSION), "text": content,
                       "payload": _base_payload(knowledge_id=f"USG-{module}-CAP-{hashlib.sha1(cap['name'].encode()).hexdigest()[:10]}",
                                                module=module, capability=cap["name"], functionality=join(f["screenId"] for f in feats),
                                                capability_basis=",".join(cap["basis"]), capability_needs_review=cap["needsReview"],
                                                knowledge_type="UsageCapability", title=f"{cap['name']} ({catalog['moduleName']})",
                                                content=content, source_types=["Usage Catalog"],
                                                features=[f["name"] for f in feats], ready_features=cap["ready"])})

    for f in catalog["features"]:
        verified = [c["text"] for c in f["claims"] if c["status"] == "VERIFIED"]
        review = [c["text"] for c in f["claims"] if c["status"] != "VERIFIED"]
        summary = (f.get("summary") or {}).get("text", "")
        parts = [f"{f['name']} ({f['category']}) in {catalog['moduleName']}, capability {f['capability']['name']}."]
        if verified:
            parts.append("Verified: " + " ".join(verified))
        if summary:
            parts.append("Description (AI draft): " + summary)
        content = " ".join(parts)
        evidence = [{"source": e.get("source", ""), "detail": e.get("detail", "")}
                    for c in f["claims"] if c["status"] == "VERIFIED" for e in c.get("evidence", [])][:12]
        points.append({"id": _uuid("usage", module, "feature", f["screenId"], INGEST_VERSION), "text": content,
                       "payload": _base_payload(
                           knowledge_id=f"USG-{f['screenId']}", module=module, capability=f["capability"]["name"],
                           functionality=func.get(f["screenId"], ""),
                           capability_basis=f["capability"]["basis"], capability_needs_review=f["capability"]["needsReview"],
                           feature=f["name"], screen_id=f["screenId"], screen_name=f["name"], knowledge_type="UsageFeature",
                           title=f["name"], content=content, source_types=["Usage Catalog", "Source Code", "Database"],
                           evidence_ids=[e["source"] for e in evidence if e["source"]],
                           review_required=f["readiness"] == "NEEDS_EVIDENCE",
                           publication_status="PUBLISHED" if verified else "REVIEW_REQUIRED",
                           readiness=f["readiness"], category=f["category"], verified_claims=verified,
                           review_claims=review, ai_summary=summary, evidence=evidence,
                           migration_gaps=len(f.get("migrationGaps", [])), alias_screen_ids=f.get("aliasScreenIds", []))})
    return with_vectors(points, kms)


def publish_usage_catalog(kms: QdrantKMS, catalog: dict[str, Any], batch_size: int = 64,
                          functionality: dict[str, str] | None = None) -> dict[str, Any]:
    kms.ensure_collection()
    points = usage_catalog_points(catalog, kms, functionality)
    for i in range(0, len(points), batch_size):
        kms.client.upsert(kms.collection, points[i:i + batch_size])
    kinds: dict[str, int] = {}
    for p in points:
        kinds[p["payload"]["knowledge_type"]] = kinds.get(p["payload"]["knowledge_type"], 0) + 1
    return {"status": "PUBLISHED", "collection": kms.collection, "module": catalog["module"], "points": len(points), "byType": kinds}


# ---------------------------------------------------------------- tickets

TICKET_FIELDS = {
    "ticket_id": ("ticket_id", "ticket id", "ticket no", "ticket number", "ticketno", "issue id", "id", "#", "call number", "callnumber", "reference"),
    "title": ("title", "subject", "summary", "issue", "short description"),
    "description": ("description", "details", "reported problem", "reportedproblem", "problem", "problem description", "body", "issue description"),
    "module": ("module", "product module", "component", "area", "product", "application"),
    "status": ("status", "state", "call status", "callstatus"),
    "priority": ("priority", "severity", "call priority", "callpriorityid"),
    "category": ("category", "type", "tracker", "issue type", "nature", "call type", "callnatureid", "calltypeid"),
    "created_at": ("created_at", "created", "created on", "created date", "date", "reported on", "call date", "calldate", "opened"),
    "resolution": ("resolution", "action taken", "actiontaken", "solution", "fix", "resolution notes"),
}
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?<!\w)\+?\d[\d\s\-()]{7,}\d(?!\w)")
SQL_FORBIDDEN_RE = re.compile(r"\b(insert|update|delete|merge|drop|alter|create|truncate|exec|execute|grant|revoke|into)\b", re.I)


def redact(text: str) -> str:
    return PHONE_RE.sub("[phone]", EMAIL_RE.sub("[email]", text or ""))


def _key(s: str) -> str:
    return re.sub(r"[^a-z0-9#]+", " ", str(s).lower()).strip()


def map_columns(columns: Iterable[str], explicit: dict[str, str] | None = None) -> dict[str, str]:
    """Return {pie_field: source_column} by explicit mapping or synonym match."""
    cols = list(columns)
    out: dict[str, str] = {}
    for field, col in (explicit or {}).items():
        if col in cols:
            out[field] = col
    by_key = {_key(c): c for c in cols}
    for field, synonyms in TICKET_FIELDS.items():
        if field in out:
            continue
        for syn in synonyms:
            if _key(syn) in by_key and by_key[_key(syn)] not in out.values():
                out[field] = by_key[_key(syn)]
                break
    return out


def read_ticket_file(path: str | Path) -> tuple[list[str], list[dict[str, Any]]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        with p.open(newline="", encoding="utf-8-sig", errors="replace") as fh:
            reader = csv.DictReader(fh)
            return list(reader.fieldnames or []), list(reader)
    if ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(next(rows, []))]
        return header, [dict(zip(header, r)) for r in rows if any(v is not None for v in r)]
    if ext == ".json":
        data = json.loads(p.read_text(encoding="utf-8"))
        rows = data if isinstance(data, list) else data.get("tickets", data.get("issues", []))
        header = sorted({k for r in rows for k in r})
        return header, rows
    raise ValueError(f"Unsupported ticket file type: {ext} (use .csv, .xlsx or .json)")


def validate_select(sql: str) -> str:
    """Allow one read-only SELECT/WITH statement only."""
    body = re.sub(r"--[^\n]*|/\*.*?\*/", " ", sql, flags=re.S).strip().rstrip(";").strip()
    no_strings = re.sub(r"'(?:''|[^'])*'", "''", body)
    if not re.match(r"^(select|with)\b", no_strings, re.I):
        raise ValueError("The ticket query must be a single SELECT (or WITH ... SELECT) statement.")
    if ";" in no_strings or SQL_FORBIDDEN_RE.search(no_strings):
        raise ValueError("The ticket query may only read data: one statement, no INSERT/UPDATE/DELETE/EXEC/INTO/DDL.")
    return body


def read_ticket_sql(sql: str, database_url: str | None = None, limit: int = 50000) -> tuple[list[str], list[dict[str, Any]]]:
    import pyodbc
    from sqlalchemy.engine import make_url
    from .db_discovery import connection_string, load_dotenv
    load_dotenv()
    query = validate_select(sql)
    u = make_url(database_url or connection_string())
    server = (u.host or "") + (f",{u.port}" if u.port else "")
    cn = pyodbc.connect(f"DRIVER={{{u.query.get('driver', 'ODBC Driver 18 for SQL Server')}}};SERVER={server};DATABASE={u.database};"
                        f"UID={u.username};PWD={u.password or ''};TrustServerCertificate=yes;Encrypt=no;ApplicationIntent=ReadOnly",
                        timeout=15, readonly=True)
    try:
        cur = cn.cursor()
        cur.execute(query)
        header = [d[0] for d in cur.description]
        rows = [dict(zip(header, r)) for r in cur.fetchmany(limit)]
    finally:
        cn.close()
    return header, rows


def _as_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return re.sub(r"\s+", " ", str(v)).strip()


def normalize_tickets(header: list[str], rows: list[dict[str, Any]], source: str,
                      mapping: dict[str, str] | None = None, known_modules: set[str] | None = None) -> tuple[list[dict[str, Any]], dict[str, str]]:
    cols = map_columns(header, mapping)
    if "title" not in cols and "description" not in cols:
        raise ValueError(f"Could not find a title/subject or description column in: {header}. Pass a --mapping file.")
    known = {re.sub(r"[^a-z0-9]", "", m.lower()): m for m in (known_modules or set())}
    tickets = []
    for i, r in enumerate(rows):
        get = lambda f: _as_text(r.get(cols[f])) if f in cols else ""
        title, desc = redact(get("title")), redact(get("description"))
        if not (title or desc):
            continue
        raw_module = get("module")
        module = known.get(re.sub(r"[^a-z0-9]", "", raw_module.lower()), raw_module.lower())
        tickets.append({
            "ticket_id": get("ticket_id") or f"{source}-{i + 1}",
            "title": title or desc[:120],
            "description": desc[:1500],
            "module": module,
            "status": get("status"), "priority": get("priority"), "category": get("category"),
            "created_at": get("created_at")[:32], "resolution": redact(get("resolution"))[:800],
            "source": source,
        })
    return tickets, cols


def ticket_points(tickets: list[dict[str, Any]], kms: QdrantKMS) -> list[dict[str, Any]]:
    points = []
    for t in tickets:
        content = " ".join(x for x in (t["title"], t["description"],
                                       f"Resolution: {t['resolution']}" if t["resolution"] else "") if x)
        points.append({"id": _uuid("ticket", t["source"], t["ticket_id"]), "text": content + " " + t["module"] + " " + t["category"],
                       "payload": _base_payload(knowledge_id=f"TKT-{t['source']}-{t['ticket_id']}", module=t["module"], knowledge_type="Ticket",
                                                title=t["title"], content=content[:2500], source_types=["Ticket"],
                                                ticket_id=t["ticket_id"], ticket_status=t["status"], ticket_priority=t["priority"],
                                                ticket_category=t["category"], ticket_created_at=t["created_at"], ticket_source=t["source"],
                                                audience="internal")})
    return with_vectors(points, kms)


def publish_tickets(kms: QdrantKMS, tickets: list[dict[str, Any]], batch_size: int = 128) -> dict[str, Any]:
    kms.ensure_collection()
    points = ticket_points(tickets, kms)
    for i in range(0, len(points), batch_size):
        kms.client.upsert(kms.collection, points[i:i + batch_size])
    return {"status": "PUBLISHED", "collection": kms.collection, "tickets": len(points)}


def ticket_summary(tickets: list[dict[str, Any]], cols: dict[str, str]) -> dict[str, Any]:
    from collections import Counter
    return {"tickets": len(tickets), "columnMapping": cols,
            "byModule": dict(Counter(t["module"] or "(none)" for t in tickets).most_common(15)),
            "byStatus": dict(Counter(t["status"] or "(none)" for t in tickets).most_common(10)),
            "byPriority": dict(Counter(t["priority"] or "(none)" for t in tickets).most_common(10))}


# ---------------------------------------------------------------- field knowledge (v0.20)

def field_points(dictionary: dict[str, Any], kms: QdrantKMS, context: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    """One FieldKnowledge point per field, carrying the full field record and the screen's business
    context (functionality, capability) so every point holds the whole chain down to the field."""
    from .field_intel import field_summary_text
    points = []
    for r in dictionary["fields"]:
        content = field_summary_text(r)
        search_text = " ".join([content, r["fieldName"], r["label"], r["database"].get("column", ""), r["page"], r["section"],
                                (context or {}).get("capability", ""), (context or {}).get("functionality", "")])
        record = {k: v for k, v in r.items() if k not in {"factHash"}}
        key = r.get("fieldKey") or r["fieldName"]
        points.append({"id": _uuid("field", r["screenId"], key, INGEST_VERSION), "text": search_text,
                       "payload": _base_payload(
                           knowledge_id=f"FLD-{r['screenId']}-{key}", module=r["module"], feature=r["screen"],
                           screen_id=r["screenId"], screen_name=r["screen"], knowledge_type="FieldKnowledge",
                           title=f"{r['screen']} field: {r['label']} ({r['fieldName']})", content=content,
                           source_types=["Frontend Code", "Backend Code", "Database"], evidence_ids=[e["source"] for e in r["evidence"]][:10],
                           generation_method="deterministic+ai" if r.get("explanation") else "deterministic",
                           field_name=r["fieldName"], field_label=r["label"], field_page=r["page"], field_section=r["section"],
                           field_mandatory=r["mandatory"]["required"],
                           field_column=r["database"].get("column") or (r["database"].get("storedVia") or {}).get("column", ""),
                           field_stored_via=(r["database"].get("storedVia") or {}).get("field", ""),
                           field=record, **_context(context))})
    return with_vectors(points, kms)


def section_points(dictionary: dict[str, Any], kms: QdrantKMS, context: dict[str, Any] | None = None) -> list[dict[str, Any]]:
    """One SectionKnowledge point per screen tab/section: its fields, mandatory ones and rules."""
    groups: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for r in dictionary["fields"]:
        if r.get("visibleOnScreen"):
            groups.setdefault((r.get("page") or "", r.get("section") or ""), []).append(r)
    points = []
    for (page, section), rs in groups.items():
        where = " > ".join(x for x in (page, section) if x) or "main form"
        mand = [r["label"] for r in rs if r["mandatory"]["required"]]
        calc = [f"{r['label']} = {r['calculation']['formula']}" for r in rs if r.get("calculation")]
        rules = sorted({v["rule"] for r in rs for v in r["validations"] if v["layer"] == "backend"})
        tables = sorted({r["database"]["table"] for r in rs if r["database"].get("table")})
        content = (f"{dictionary['screen']} ({dictionary['module']}) section {where}: {len(rs)} fields: "
                   + "; ".join(f"{r['label']} ({(r.get('type') or {}).get('ui') or r['type']['form']}{', mandatory' if r['mandatory']['required'] else ''})" for r in rs)
                   + (f". Mandatory: {', '.join(mand)}" if mand else "") + (f". Calculated: {'; '.join(calc)}" if calc else "")
                   + (f". Backend validations: {'; '.join(rules)}" if rules else "") + (f". Stored in: {', '.join(tables)}" if tables else "") + ".")
        points.append({"id": _uuid("section", dictionary["screenId"], page, section, INGEST_VERSION), "text": content,
                       "payload": _base_payload(
                           knowledge_id=f"SEC-{dictionary['screenId']}-{hashlib.sha1(where.encode()).hexdigest()[:8]}",
                           module=dictionary["module"], feature=dictionary["screen"], screen_id=dictionary["screenId"],
                           screen_name=dictionary["screen"], knowledge_type="SectionKnowledge", title=f"{dictionary['screen']} section: {where}",
                           content=content[:3000], source_types=["Frontend Code"], section_page=page, section_name=section,
                           section_fields=[r["fieldName"] for r in rs], section_mandatory=mand, **_context(context))})
    return with_vectors(points, kms)


def publish_field_dictionary(kms: QdrantKMS, dictionary: dict[str, Any], batch_size: int = 32,
                             context: dict[str, Any] | None = None) -> dict[str, Any]:
    """`context`: the screen's functionality/capability (see screen_contexts)."""
    kms.ensure_collection()
    points = field_points(dictionary, kms, context)
    sections = section_points(dictionary, kms, context)
    for pts in (points, sections):
        for i in range(0, len(pts), batch_size):
            kms.client.upsert(kms.collection, pts[i:i + batch_size])
    return {"status": "PUBLISHED", "collection": kms.collection, "screen": dictionary["screen"], "fields": len(points),
            "sections": len(sections), "explained": sum(1 for r in dictionary["fields"] if r.get("explanation"))}

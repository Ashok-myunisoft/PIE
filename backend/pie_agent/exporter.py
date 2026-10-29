from __future__ import annotations
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from .models import PIEOutput

def export_json(out: PIEOutput, path: Path):
    path.write_text(json.dumps(out.model_dump(mode="json"),indent=2,ensure_ascii=False),encoding="utf-8")

def export_traceability(out: PIEOutput, path: Path):
    path.write_text(json.dumps(out.traceability.model_dump(mode="json"),indent=2,ensure_ascii=False),encoding="utf-8")

def export_review(out: PIEOutput, path: Path):
    path.write_text(json.dumps([x.model_dump(mode="json") for x in out.reviewItems],indent=2,ensure_ascii=False),encoding="utf-8")

def export_validation(report, path: Path):
    path.write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding="utf-8")

def export_markdown(out: PIEOutput, path: Path):
    lines=[f"# {out.module.ModuleName}", "", f"**Functional area:** {out.module.FunctionalArea}", "",
           "## Config Items"]
    for x in out.configItems:
        lines.append(f"- **{x.ItemName}** — {x.Description} — mandatory={x.IsMandatory}, hours={x.EstimatedHours}, review={x.needsReview}")
    lines += ["","## Variants"]
    for x in out.variants:
        lines.append(f"- **{x.VariantName}** — {x.PlainLanguageDesc} — review={x.needsReview}")
    lines += ["","## Training Topics"]
    for x in out.trainingTopics:
        lines.append(f"- **{x.TopicName}** — {x.MenuPath}")
    lines += ["","## UAT Scenarios"]
    for x in out.uatScenarios:
        lines.append(f"- **{x.ScenarioTitle}** — {x.ExpectedResult}")
    path.write_text("\n".join(lines),encoding="utf-8")

def export_excel(out: PIEOutput, path: Path, prompt: str):
    wb=Workbook()
    ws=wb.active; ws.title="PIE Output"
    headers=["S.No","Content Discovery","Sub Category","Input","Output","Summary","Prompt"]
    ws.append(headers)
    ws.append([
        1,"Multi-source synthesis","Leave Application / Leave Request",
        "Relevant screen + frontend + backend + DB + API + manual/training evidence",
        "Generated PIE JSON",
        "Business Product Intelligence Summary",
        prompt
    ])
    for c,w in enumerate([8,24,34,55,28,34,90],1): ws.column_dimensions[get_column_letter(c)].width=w
    ws.freeze_panes="A2"
    ws2=wb.create_sheet("Business Summary")
    if out.businessSummary:
        cols=list(out.businessSummary[0].keys()); ws2.append(cols)
        for row in out.businessSummary: ws2.append([row.get(c,"") for c in cols])
        for i in range(1,len(cols)+1): ws2.column_dimensions[get_column_letter(i)].width=28
    wb.save(path)
